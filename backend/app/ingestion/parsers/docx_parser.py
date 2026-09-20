"""DOCX parser - preserves the interleaved order of paragraphs and tables."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.parsers import register
from app.ingestion.types import NormalizedDocument, clean_text, looks_like_heading

_HEADING_STYLE_PREFIX = "Heading"
_PAGE_BREAK_TAG = qn("w:br")


def _iter_body_items(document: DocxDocument):
    """Yield Paragraph and Table objects in true document order.

    python-docx exposes `.paragraphs` and `.tables` as two separate lists, which
    loses their relative order. For a document like "intro, table, explanation of
    that table" the order matters enormously for retrieval quality, so we walk the
    underlying XML body instead.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _count_page_breaks(paragraph: Paragraph) -> int:
    """Count explicit page breaks so we can approximate page numbers.

    Word does not store rendered page numbers in the file - pagination depends on
    the printer and the fonts installed. We therefore report an ESTIMATE derived
    from hard page breaks, and label it as such in the citation UI. Being honest
    about an approximation is far better than showing a confident wrong page.
    """
    count = 0
    for br in paragraph._p.iter(_PAGE_BREAK_TAG):
        if br.get(qn("w:type")) == "page":
            count += 1
    return count


@register(".docx", label="Word document")
def parse_docx(path: Path, document_id: int, document_name: str) -> NormalizedDocument:
    doc = NormalizedDocument(
        document_id=document_id, document_name=document_name, file_type="docx"
    )

    document = DocxDocument(str(path))

    core = document.core_properties
    doc.doc_metadata = {
        "title": core.title or "",
        "author": core.author or "",
        "subject": core.subject or "",
        "created": str(core.created) if core.created else "",
        "modified": str(core.modified) if core.modified else "",
        "page_numbers_are_estimated": True,
    }

    page = 1
    current_section = ""
    table_index = 0

    for item in _iter_body_items(document):
        if isinstance(item, Paragraph):
            page += _count_page_breaks(item)
            text = clean_text(item.text)
            if not text:
                continue

            style_name = (item.style.name if item.style is not None else "") or ""
            if style_name.startswith(_HEADING_STYLE_PREFIX) or looks_like_heading(text):
                current_section = text[:150]
                block_type = "heading"
            elif style_name.startswith("List") or text.lstrip().startswith(("•", "-", "*", "–")):
                block_type = "bullet"
            else:
                block_type = "paragraph"

            doc.add_block(
                text,
                block_type=block_type,
                page_number=page,
                page_end=page,
                section=current_section,
                style=style_name,
            )

        elif isinstance(item, Table):
            table_index += 1
            doc.add_block(*_render_table(item, table_index, page, current_section))

    if page > 1:
        doc.page_count = page
    return doc


def _render_table(table: Table, table_index: int, page: int, section: str) -> tuple:
    """Flatten a table into header-labelled rows.

    A raw row like `["Q3", "12", "9"]` embeds poorly and cites badly. Rendering it
    as `Quarter: Q3 | Revenue: 12 | Profit: 9` makes each row self-describing, so
    it retrieves on the *meaning* of the column names and reads sensibly when the
    user expands the citation.
    """
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([clean_text(cell.text) for cell in row.cells])

    if not rows:
        return (
            f"[Empty table {table_index}]",
            "table",
            {"page_number": page, "section": section, "table_index": table_index},
        )

    header = rows[0]
    body = rows[1:] or rows[1:]

    lines: list[str] = []
    for r_offset, row in enumerate(body, start=2):
        pairs = []
        for c_index, cell in enumerate(row):
            if not cell:
                continue
            label = header[c_index] if c_index < len(header) and header[c_index] else f"Column {c_index + 1}"
            pairs.append(f"{label}: {cell}")
        if pairs:
            lines.append(" | ".join(pairs))

    text = f"Table {table_index}\n" + "\n".join(lines) if lines else f"Table {table_index}"

    return (
        text,
        "table",
        {
            "page_number": page,
            "section": section,
            "table_index": table_index,
            "row_start": 2,
            "row_end": len(rows),
        },
    )
